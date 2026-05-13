# 文件位置: contract-review/data/samples/generate_samples.py
# 功能说明: 示例合同生成器（开发用，非运行时依赖）。
#          生成两份覆盖"标准 + 风险"用例的合同：软件采购合同（知识产权/付款风险）、
#          大额销售合同（违约责任不对等/管辖/违约金/保密期限缺失），各输出 DOCX 与 PDF 两种格式，
#          用于验证解析引擎对 DOCX / PDF 的兼容，以及审查引擎对高风险/中风险条款的命中。
from __future__ import annotations

from pathlib import Path

from docx import Document

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

BASE = Path(__file__).resolve().parent


def _docx_body() -> dict:
    """返回 {文件名前缀: 章节列表}。每章为 (标题, 正文)。"""
    contracts = {}

    # 1) 软件采购合同（付款前置验收缺失 + 违约责任不对等）
    contracts["软件采购合同_知识产权与付款风险"] = [
        ("合同名称", "软件采购合同"),
        ("签约主体", "甲方（采购方）：星辰科技有限公司，统一社会信用代码：91110108MA01ABCD23；"
                     "乙方（供应商）：云栖软件股份有限公司，统一社会信用代码：91330106MA02EFGH45。"),
        ("合同金额与期限", "本合同金额：¥1,200,000 元（人民币）。履行期限：自合同签订之日起 12 个月内。"),
        ("第一条 标的物", "乙方向甲方提供“企业知识库管理系统”软件一套及为期一年的技术支持服务，"
                        "软件版本以乙方交付为准。"),
        ("第二条 价款与付款", "本合同采用货到付款方式，甲方应在货物（软件授权）到货后 3 个工作日内支付全部合同款项，"
                          "乙方收款后开具发票。"),
        ("第三条 知识产权", "本软件及相关全部知识产权（包括但不限于源代码、文档、算法）均归乙方所有，"
                        "甲方仅享有非独占的使用许可，不得反向工程或披露。"),
        ("第四条 验收标准", "甲方应在收到软件后 5 个工作日内完成功能核对，逾期未提出书面异议视为验收通过。"),
        ("第五条 违约责任", "因本软件任何问题导致甲方损失的，由甲方独自承担无限责任；"
                        "乙方不承担任何连带或间接责任。"),
        ("第六条 保密义务", "双方对合作中知悉的对方商业秘密予以保密。"),
        ("第七条 数据安全", "乙方在处理甲方数据时应采取合理安全措施，防止数据泄露。"),
        ("第八条 争议解决管辖", "因本合同发生争议，提交合同签订地有管辖权的人民法院诉讼解决。"),
    ]

    # 2) 大额销售合同（违约责任不对等 + 管辖原告方 + 违约金超限 + 保密无期限）
    contracts["大额销售合同_违约责任不对等"] = [
        ("合同名称", "产品销售合同"),
        ("签约主体", "甲方（买方）：华夏制造集团有限公司，统一社会信用代码：91440300MA03WJKL78；"
                     "乙方（卖方）：锐进设备有限公司，统一社会信用代码：91310115MA04MNBP90。"),
        ("合同金额与期限", "本合同总金额：¥8,600,000 元（人民币）。履行期限：自合同生效之日起 18 个月。"),
        ("第一条 标的物", "乙方向甲方供应自动化生产线设备一套，含安装调试与操作培训。"),
        ("第二条 价款与付款", "甲方分三期支付：预付款 30%，到货付 60%，余款 10% 于质保期满支付；"
                          "每期付款以乙方书面通知为准，甲方应及时支付。"),
        ("第三条 验收标准", "设备到厂后由乙方主导验收，甲方配合；验收报告由乙方出具。"),
        ("第四条 违约责任", "若乙方违约，应承担由此给甲方造成的一切损失；若甲方违约，相关责任由甲方自行承担，"
                        "乙方不追究。违约金比例为合同总额的 35%。"),
        ("第五条 保密义务", "双方对交易中获知的对方技术资料与商业信息负有保密义务。"),
        ("第六条 数据安全", "双方在数据交互中应遵守相关数据保护要求。"),
        ("第七条 争议解决管辖", "因本合同引起的争议，由甲方（原告）住所地人民法院管辖。"),
    ]
    return contracts


def build_docx(name: str, sections: list, out_dir: Path):
    doc = Document()
    doc.add_heading(name, level=0)
    for title, body in sections:
        doc.add_heading(title, level=1)
        doc.add_paragraph(body)
    path = out_dir / f"{name}.docx"
    doc.save(str(path))
    return path


def build_pdf(name: str, sections: list, out_dir: Path):
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    style = ParagraphStyle("cn", fontName="STSong-Light", fontSize=11, leading=18)
    title_style = ParagraphStyle("cnt", fontName="STSong-Light", fontSize=15,
                                 leading=24, spaceAfter=10)
    path = out_dir / f"{name}.pdf"
    doc = SimpleDocTemplate(str(path), pagesize=A4,
                            leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48)
    flow = [Paragraph(name, title_style)]
    for title, body in sections:
        flow.append(Paragraph(f"<b>{title}</b>", style))
        flow.append(Paragraph(body, style))
        flow.append(Spacer(1, 6))
    doc.build(flow)
    return path


def main():
    out = BASE
    contracts = _docx_body()
    for name, sections in contracts.items():
        p1 = build_docx(name, sections, out)
        p2 = build_pdf(name, sections, out)
        print("生成:", p1.name, "/", p2.name)


if __name__ == "__main__":
    main()
