# 文件位置: contract-review/tests/regression.py
# 功能说明: 端到端回归自检脚本。
#          覆盖四条关键路径：
#            ① 非合同文档（技术笔记/面试题）必须被「合同闸门」拦下，不产生任何风险发现；
#            ② 真实示例合同必须正常完成审查，且风险高亮的字符偏移与原文严格对齐；
#            ③ 多张照片识别出的多页正文必须按页码合并、空白页不参与审查，
#               且合并后的高亮偏移仍然逐字对齐（分页符不能把偏移搞错位）；
#            ④ 审查失败的任务必须经「条件门禁 → 人工确认」才外发，且驳回留理由、
#               已推送不重复、重跑作废旧记录（这几条是"失败通知"不变成噪音源的前提）。
#          直接运行：python tests/regression.py
from __future__ import annotations

import sys
from pathlib import Path

# 允许以脚本方式直接运行（把项目根目录加入模块搜索路径）
ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from config import cfg  # noqa: E402
from core.review_pipeline import ReviewPipeline  # noqa: E402
from core.models import TaskStatus  # noqa: E402
from services.ocr_service import (  # noqa: E402
    ImageQuality,
    OCRResult,
    PageOCR,
    assemble_batch,
)

FIXTURES = ROOT / "tests" / "fixtures"
SAMPLES = ROOT / "data" / "samples"

# 非合同样例正文：仿照"把技术笔记误传进审查系统"的场景
NON_CONTRACT_PARAS = [
    "智能体核心技术",
    "Python基础",
    "数据结构相关",
    "**list 和 tuple 有什么区别？",
    "列表（list）是可变的，可执行添加、删除或修改元素等操作；元组（tuple）不可变，创建后内容不能更改。"
    "但元组可以被哈希，可用作字典的键，但列表不行。",
    "**闭包相关",
    "什么是闭包？请简述其定义和形成条件。",
    "闭包是一种特殊的函数嵌套结构，指在外部函数内部定义的内部函数，引用了外部函数的变量（非全局变量），"
    "且外部函数返回该内部函数。",
    "形成闭包需满足三个条件：",
    "存在函数嵌套（内部函数定义在外部函数中）；",
    "内部函数引用了外部函数的变量；",
    "外部函数返回内部函数的引用。",
    "这种机制依赖 Python 的作用域链：内部函数在查找变量时，会先在自身作用域查找，"
    "若未找到则向上级（外部函数）作用域查找，闭包通过这种方式查找保留住对外部变量的引用。",
    "闭包中如何修改外部函数的变量？nonlocal 关键字的作用是什么？",
    "在闭包中，内部函数默认只能读取外部函数的变量，若直接修改会被 Python 判定为“定义局部变量”而报错。",
]


def make_non_contract_docx(path: Path) -> None:
    """生成一份非合同文档（技术笔记）用于闸门回归测试。"""
    from docx import Document

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    doc.add_heading("Python 基础知识点整理", level=1)
    for p in NON_CONTRACT_PARAS:
        doc.add_paragraph(p)
    doc.save(str(path))


def check_offsets(task) -> bool:
    """校验所有风险发现的高亮偏移与原文逐字一致。"""
    ok = True
    for f in task.findings:
        if f.start_offset is None or f.end_offset is None:
            continue
        seg = task.full_text[f.start_offset:f.end_offset]
        if seg != f.quote:
            ok = False
            print(f"    [偏移错位] {f.finding_id}: {seg!r} != {f.quote!r}")
    return ok


def main() -> int:
    pipe = ReviewPipeline()
    failures = []

    # ---- 用例一：非合同文档必须被拦下 ----
    nc_path = FIXTURES / "非合同样例_Python技术笔记.docx"
    make_non_contract_docx(nc_path)
    t = pipe.create_task(filename=nc_path.name, contract_type="未分类")
    r = pipe.run(t.task_id, str(nc_path))
    dc = r.doc_check
    print("=" * 62)
    print(f"[用例1] {nc_path.name}")
    print(f"  状态：{r.status.value}（期望 rejected）")
    print(f"  合同判定：{dc.is_contract} | 得分 {dc.score}/{dc.threshold} | 负向特征 {len(dc.negative_hits)} 项")
    print(f"  拒绝原因：{dc.reject_reason}")
    print(f"  风险发现数：{len(r.findings)}（期望 0）")
    if r.status != TaskStatus.REJECTED:
        failures.append("非合同文档未被拦截")
    if r.findings:
        failures.append("非合同文档产生了风险发现（误报未消除）")
    if dc and dc.score >= dc.threshold:
        failures.append("非合同文档得分未低于阈值")

    # ---- 用例二：真实合同必须正常审查 ----
    for name, ctype in [("软件采购合同_知识产权与付款风险.pdf", "采购"),
                        ("大额销售合同_违约责任不对等.pdf", "销售")]:
        src = SAMPLES / name
        if not src.exists():
            print(f"  [跳过] 示例合同不存在：{name}")
            continue
        t = pipe.create_task(filename=name, contract_type=ctype)
        r = pipe.run(t.task_id, str(src))
        print("=" * 62)
        print(f"[用例2] {name}")
        print(f"  状态：{r.status.value}（期望 completed）| 综合风险：{r.report.overall_level.value if r.report else '-'}")
        print(f"  合同判定：{r.doc_check.is_contract} | 得分 {r.doc_check.score}/{r.doc_check.threshold}")
        print(f"  主体：{r.meta.party_a} / {r.meta.party_b} | 信用代码：{r.meta.credit_code} | 金额：{r.meta.amount}")
        for f in r.findings:
            print(f"    [{f.level.value:6}] {f.category} | {f.quote[:26]!r} | {f.source}")
        if r.status != TaskStatus.COMPLETED:
            failures.append(f"{name} 未完成审查")
        if not r.findings:
            failures.append(f"{name} 未识别出任何风险（漏报）")
        if not check_offsets(r):
            failures.append(f"{name} 高亮偏移与原文不一致")

    # ---- 用例三：多张照片（多页合同）的分页合并 ----
    # 不依赖任何 OCR 引擎：直接构造"逐页识别成功/失败"的结果对象，
    # 验证的是合并规则与解析器的约定（这一层出错，风险卡片的页码就会指错）。
    photo_pages = [
        "采购合同\n甲方：北京星辰科技有限公司\n乙方：上海云海设备有限公司\n合同金额：¥1,200,000 元",
        "",   # 空白页（如背面、盖章页）：不应参与审查
        "第三条 付款：合同签订后 30 日内支付全款。\n第九条 保密义务不设期限。",
    ]
    page_results = [
        PageOCR(index=i + 1, filename=f"照片_第{i+1}页.jpg", bytes=1024,
                result=OCRResult(ok=True, text=t, provider="fake", provider_label="假引擎（自检）",
                                 quality=ImageQuality()) if t.strip()
                else OCRResult(ok=False, reason="未识别出文字", hint="人工录入"))
        for i, t in enumerate(photo_pages)
    ]
    batch = assemble_batch(page_results)
    txt = FIXTURES / "多页照片合同_自检.txt"
    txt.parent.mkdir(parents=True, exist_ok=True)
    txt.write_text(batch["text"], encoding="utf-8")

    # 分页符是否被解析器认出来，直接问解析器（任务快照里不存 pages）
    from services.doc_parser import DocParser  # noqa: E402

    parsed_pages = len(DocParser().parse(str(txt))["pages"])
    t = pipe.create_task(filename="多页照片合同（拍照识别）·3页", contract_type="采购")
    r = pipe.run(t.task_id, str(txt))
    page_nos = sorted({s.page for s in (r.segments or []) if s.page})
    print("=" * 62)
    print("[用例3] 多页照片合同_自检.txt")
    print(f"  逐页识别：成功 {batch['ok_pages']}/3 页 | 未识别页：{batch['failed_pages']}（期望 [2]）")
    print(f"  合并后解析页数：{parsed_pages}（期望 3：空白页保留页次，页码才与照片序号对齐）")
    print(f"  条款页码分布：{page_nos}（期望 [1, 3]：正文在第 1、3 张照片上，第 2 张空白留位）")
    print(f"  状态：{r.status.value} | 综合风险：{r.report.overall_level.value if r.report else '-'}")
    for f in r.findings:
        print(f"    [{f.level.value:6}] {f.category} | {f.quote[:26]!r}")
    if batch["failed_pages"] != [2]:
        failures.append("空白页未被识别为“未识别”，逐页状态错误")
    if parsed_pages != 3:
        failures.append(f"多页合并后页数应为 3（空页留位），实际 {parsed_pages}（分页符可能失效）")
    # 页码必须与"照片序号"一致：第 3 张照片上的条款就该落在第 3 页，
    # 若空页被剔除则会前移成第 2 页 —— 用户照页码去翻照片就会对不上。
    if page_nos != [1, 3]:
        failures.append(f"条款页码应为 [1, 3]（与照片序号对齐），实际 {page_nos}")
    if r.status != TaskStatus.COMPLETED:
        failures.append("多页照片合同未完成审查")
    if not r.findings:
        failures.append("多页照片合同未识别出任何风险（漏报）")
    if not check_offsets(r):
        failures.append("多页合并后的高亮偏移与原文不一致")

    # ---- 用例四：失败任务的外发审核链路（条件门禁 + 人工确认）----
    # 这条链路的价值全在"分寸"上：既不能把噪音（随手传错文件）涌进钉钉群，
    # 也不能让过期失败（已重试成功）被误发。所以逐条断言这几个约束。
    from services.export_queue import export_queue  # noqa: E402
    from services.failure_gate import normalize_condition  # noqa: E402

    # 造两个失败任务：一个真解析不了（blocked），一个是非合同（rejected）
    broken_pdf = FIXTURES / "加密扫描件_无法解析.pdf"
    broken_pdf.parent.mkdir(parents=True, exist_ok=True)
    broken_pdf.write_bytes(b"%PDF-1.4\n" + b"\x00\x01\x02 not a real pdf " * 40)
    tb = pipe.create_task(filename=broken_pdf.name, contract_type="采购")
    r_blocked = pipe.run(tb.task_id, str(broken_pdf))
    tj = pipe.create_task(filename=nc_path.name, contract_type="未分类")
    r_rejected = pipe.run(tj.task_id, str(nc_path))

    # 出厂默认条件只纳入「解析受阻」：闸门拦下非合同是正常工作结果，不该打扰钉钉群
    cfg.EXPORT_CONDITION = dict(cfg.EXPORT_CONDITION,
                                kinds=["blocked"], min_retries=0,
                                extensions=[], sources=[], reason_keywords=[], dedupe=True)
    export_queue._records.clear()
    scan1 = export_queue.scan([r_blocked, r_rejected])
    recs = {x["task_id"]: x for x in export_queue.list_records()}
    b_rec = recs.get(r_blocked.task_id, {})
    j_rec = recs.get(r_rejected.task_id, {})
    print("=" * 62)
    print("[用例4] 失败外发审核链路")
    print(f"  两个失败任务：blocked={r_blocked.status.value} / rejected={r_rejected.status.value}")
    print(f"  默认条件（仅解析受阻）扫描：入队 {scan1['queued']} 条，条件拦下 {scan1['filtered']} 条")
    print(f"    解析受阻 -> {b_rec.get('status')}（期望 pending）")
    print(f"    非合同   -> {j_rec.get('status')}（期望 filtered）")
    if b_rec.get("status") != "pending":
        failures.append("解析受阻任务未进入待审核队列")
    if j_rec.get("status") != "filtered":
        failures.append("非合同被拒的任务本应被默认条件拦下（否则钉钉群会被传错文件的告警刷屏）")
    if j_rec.get("status") == "filtered" and "失败类型" not in j_rec.get("gate", {}).get("summary", ""):
        failures.append("被拦下的记录未说明是哪个条件没过（留痕不可解释）")

    # 消息正文必须入队时就冻结，且带上安全设置要求的关键词
    msg = b_rec.get("message", "")
    kw = (cfg.DINGTALK_KEYWORD or "").strip()
    print(f"  入队即冻结消息正文：{len(msg)} 字 | 含关键词「{kw}」：{kw in msg}")
    if kw and kw not in msg:
        failures.append("消息正文未包含 DINGTALK_KEYWORD，钉钉会因关键词不匹配拒收")

    # 驳回必须留理由，否则队列被人无声清空后无从回溯
    rid = b_rec.get("record_id")
    try:
        export_queue.reject(rid, "自检", "   ")
        failures.append("驳回未强制要求填写理由")
    except ValueError:
        print("  驳回空理由 -> 已拒绝（正确）")

    # 配错的 Webhook：真实失败必须被识别为 push_failed，而不是报成功
    cfg.DINGTALK_WEBHOOK = "http://127.0.0.1:1/robot/send?access_token=x"
    rec = export_queue.approve(rid, "自检")
    print(f"  通道配错时推送：状态={rec.status.value}（期望 push_failed）| {rec.push.detail[:44]}")
    if rec.status.value != "push_failed":
        failures.append("Webhook 不可达时未判定为推送失败")

    # 恢复可用通道后重试：正文必须沿用审核时那份，不重新生成
    cfg.DINGTALK_WEBHOOK = ""
    rec2 = export_queue.retry_push(rid, "自检")
    print(f"  通道恢复后重试：状态={rec2.status.value}（期望 pushed）| 尝试 {rec2.attempts} 次")
    if rec2.status.value != "pushed":
        failures.append("通道恢复后重试仍失败")
    if rec2.message != msg:
        failures.append("重试改了消息正文（审核过的内容必须原样发出）")

    # 去重：已推送的任务再扫描不该重新入队
    scan2 = export_queue.scan([r_blocked, r_rejected])
    print(f"  再扫描（已有已推送记录）：入队 {scan2['queued']} 条（期望 0）")
    if scan2["queued"] != 0:
        failures.append("已推送的任务被重复入队（去重失效）")

    # 条件放宽后，之前被拦下的非合同应能入队（证明条件是"可调的"而非写死）
    cfg.EXPORT_CONDITION = dict(cfg.EXPORT_CONDITION, kinds=["blocked", "rejected"])
    scan3 = export_queue.scan([r_blocked, r_rejected])
    print(f"  条件加入「非合同被拒」后：入队 {scan3['queued']} 条（期望 1）")
    if scan3["queued"] != 1:
        failures.append("放宽条件后非合同任务仍未入队（条件未生效）")

    # 任务重新审查 -> 旧待审记录必须作废，否则审核人一点就把过期失败信息发进群
    left = [x for x in export_queue.list_records("pending") if x["task_id"] == r_rejected.task_id]
    if left:
        pipe.run(r_rejected.task_id, str(nc_path))
        after = [x for x in export_queue.list_records() if x["record_id"] == left[0]["record_id"]][0]
        print(f"  任务重跑后旧待审记录：{after['status']}（期望 superseded）")
        if after["status"] != "superseded":
            failures.append("任务重新审查后，旧待审记录未作废（过期失败提醒仍可能被发出）")
    else:
        failures.append("条件放宽后未找到待审核记录，无法验证重跑作废")

    # 收尾：把外发配置恢复到出厂状态，避免污染后续手工演示
    cfg.DINGTALK_WEBHOOK = ""
    cfg.EXPORT_CONDITION = dict(cfg.EXPORT_CONDITION, kinds=["blocked"])
    assert normalize_condition(cfg.EXPORT_CONDITION).kinds == ["blocked"]

    print("=" * 62)
    if failures:
        print("自检结果：失败")
        for x in failures:
            print("  - " + x)
        return 1
    print("自检结果：全部通过")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
